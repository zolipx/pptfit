"""Every guarantee must catch its own kind of damage, and the original must pass itself.

Runs on the synthetic deck only (bench/synth.py), never on real decks.
    python -m pytest -q tests
"""
import copy, json, os, sys, tempfile

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "bench"))
import synth  # noqa: E402
from pptfit import extract as X, layout as L, render as R, textcheck as T, units as U  # noqa: E402
from pptfit.ops import Deck  # noqa: E402

WORK = tempfile.mkdtemp(prefix="pptfit_test_")


@pytest.fixture(scope="module")
def deck():
    path = synth.build(WORK)
    tr, missing = synth.translations(path)
    assert not missing
    o = X.extract(path)
    pdf = R.to_pdf(path)
    cons = {s["index"]: L.constraints(s["shapes"], R.page_words(pdf, s["index"] - 1, o["slide_w"], o["slide_h"]),
                                      o["slide_w"], o["slide_h"]) for s in o["slides"]}
    return {"path": path, "tr": tr, "orig": o, "cons": cons}


def codes(deck, path, slide):
    o = deck["orig"]
    pdf = R.render_slide(path, slide)
    cand = X.extract(path)["slides"][slide - 1]["shapes"]
    v, _ = L.check(deck["cons"][slide], cand, R.page_words(pdf, 0, o["slide_w"], o["slide_h"]))
    return {x["code"] for x in v}


def mutate(deck, name, fn):
    d = Deck(deck["path"])
    fn(d)
    out = os.path.join(WORK, name + ".pptx")
    d.save(out)
    return out


def set_text(d, eid, text):
    tx = d.idx[eid]._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}txBody")
    U.set_paragraph_text(tx.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}p")[0], text)


def ids(deck, slide, pred):
    return [e["id"] for e in deck["orig"]["slides"][slide - 1]["shapes"] if pred(e)]


def test_original_passes_itself(deck):
    for s in deck["orig"]["slides"]:
        assert codes(deck, deck["path"], s["index"]) == set(), s["index"]


def test_ink_out_of_box(deck):
    box = deck["cons"][1]["groups"][0][0]                     # a flowchart box
    p = mutate(deck, "ink_out", lambda d: set_text(d, box, "Customer orders at the kiosk and pays for every cup"))
    assert "INK_OUT" in codes(deck, p, 1)


def test_title_wraps(deck):
    t = deck["cons"][2]["titles"][0]
    p = mutate(deck, "lines", lambda d: set_text(d, t, "Highlights of the new flagship store"))
    assert codes(deck, p, 2) & {"LINES", "INK_HIT"}


def test_group_size(deck):
    g = deck["cons"][1]["groups"][0]
    p = mutate(deck, "group", lambda d: d.scale_font(g[0], 0.8, 18))
    assert "GROUP_SIZE" in codes(deck, p, 1)


def test_picture_moved(deck):
    pic = ids(deck, 3, lambda e: e["kind"] == "picture" and not e.get("inherited"))[0]
    p = mutate(deck, "moved", lambda d: d.move(pic, 10, 0))
    assert "MOVED" in codes(deck, p, 3)


def test_text_into_footer(deck):
    texts = [e for e in deck["orig"]["slides"][2]["shapes"] if e["kind"] == "text"]
    body = max(texts, key=lambda e: e["bbox"][3])["id"]       # the tall body column above the footer
    p = mutate(deck, "footer", lambda d: set_text(d, body, "\n".join(["Line %d of a very long description" % i for i in range(30)])))
    assert codes(deck, p, 3) & {"INK_HIT", "INK_BG", "OFF_SLIDE"}


def test_box_hits_neighbour(deck):
    a, b = deck["cons"][1]["groups"][0][:2]
    ea = next(e for e in deck["orig"]["slides"][0]["shapes"] if e["id"] == a)
    eb = next(e for e in deck["orig"]["slides"][0]["shapes"] if e["id"] == b)
    p = mutate(deck, "hit", lambda d: d.move(a, eb["bbox"][0] - ea["bbox"][0], eb["bbox"][1] - ea["bbox"][1]))
    assert codes(deck, p, 1) & {"BOX_HIT", "INK_HIT"}


def test_shrink_on_overflow_is_seen(deck):
    box = ids(deck, 6, lambda e: e.get("autofit") == "norm")[0]
    p = mutate(deck, "autofit", lambda d: set_text(d, box, "Opening hours 7:00 to 22:00 with 36 seats " * 12))
    assert codes(deck, p, 6) & {"FONT_FLOOR", "INK_HIT", "INK_BG", "OFF_SLIDE"}


def test_accuracy_checks(deck):
    u = next(x for x in deck["tr"] if len(x["concepts"]) >= 2)
    assert not T.check_unit(u, u["full"])
    dropped = u["full"].lower().replace(u["concepts"][0]["en"].lower(), "")
    assert any(v["code"] == "CONCEPT" for v in T.check_unit(u, dropped))
    assert any(v["code"] == "LEFTOVER" for v in T.check_unit(u, u["src"]))
    assert any(v["code"] == "DUP_WORD" for v in T.check_unit(u, "cup cup " + u["full"]))


# ---------- mechanisms that need no render ----------
def _tip(d, eid):
    from pptfit.ops import WEDGE
    prst, _, names, vals = d._callout(eid)
    x, y, w, h = d.geometry(eid)
    if prst in WEDGE:
        return x + w / 2 + vals["adj1"] * w / 1e5, y + h / 2 + vals["adj2"] * h / 1e5
    return x + vals[names[-1]] * w / 1e5, y + vals[names[-2]] * h / 1e5


def test_callout_tip_stays():
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Pt
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    a = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT, Pt(50), Pt(50), Pt(80), Pt(20))
    a.adjustments[0], a.adjustments[1] = 0.29, 4.5            # tip far below the box, like a label on a photo
    b = s.shapes.add_shape(MSO_SHAPE.LINE_CALLOUT_1, Pt(300), Pt(50), Pt(80), Pt(20))  # default adjust values
    path = os.path.join(WORK, "callout.pptx")
    prs.save(path)
    d = Deck(path)
    for sh in (a, b):
        eid = "s1_%d" % sh.shape_id
        before = _tip(d, eid)
        d.resize(eid, 0, 25)
        d.move(eid, 10, -5)
        after = _tip(d, eid)
        assert max(abs(p - q) for p, q in zip(before, after)) < 6350, (eid, before, after)  # within 0.5 pt


def _el(eid, kind, bbox, **kw):
    e = {"id": eid, "kind": kind, "bbox": bbox, "text": ""}
    e.update(kw)
    return e


def _w(text, box, block):
    return {"text": text, "box": box, "block": block, "size": 14}


def test_line_and_ragged_text_obstacles():
    els = [_el("t", "text", [60, 17, 360, 29], text="标题", is_title=True),
           _el("ln", "line", [148, 34, 549, 0], visible_line=True),   # a rule after the title: zero height
           _el("logo", "picture", [640, 327, 77, 52]),
           _el("b", "text", [428, 55, 310, 330], text="长 短")]
    orig = [_w("标题", [62, 20, 60, 25], 0), _w("长", [430, 60, 288, 15], 1), _w("短", [430, 318, 200, 15], 2)]
    con = L.constraints(els, orig, 720, 405)                   # the body's union box grazes the logo, no word does
    assert L.check(con, els, orig)[0] == []
    cand = [_w("Title", [62, 20, 70, 25], 0), _w("words", [140, 20, 60, 25], 0),
            orig[1], _w("short", [430, 318, 200, 15], 2), _w("end", [645, 333, 40, 15], 2)]
    hits = {(x["id"], x.get("other")) for x in L.check(con, els, cand)[0] if x["code"] == "INK_HIT"}
    assert ("t", "ln") in hits and ("b", "logo") in hits, hits
    # the original title's glyph box grazes the start of the rule: that licenses the graze, not more
    graze = [_w("标题", [62, 20, 90, 25], 0)] + orig[1:]
    con = L.constraints(els, graze, 720, 405)
    assert L.check(con, els, graze)[0] == []
    longer = [_w("Plan", [62, 20, 60, 25], 0), _w("layout", [126, 20, 70, 25], 0)] + orig[1:]
    assert ("t", "ln") in {(x["id"], x.get("other")) for x in L.check(con, els, longer)[0] if x["code"] == "INK_HIT"}
    # a caption that sat half on a photo may keep sitting on it
    photo = [_el("ph", "picture", [100, 200, 200, 100]), _el("cap", "text", [250, 280, 200, 30], text="说明")]
    cap0 = [_w("说明", [250, 285, 100, 20], 0)]
    con = L.constraints(photo, cap0, 720, 405)
    cap1 = [_w("Caption", [250, 285, 80, 20], 0), _w("text", [335, 285, 50, 20], 0)]
    assert not [x for x in L.check(con, photo, cap1)[0] if x["code"] == "INK_HIT"]


def test_chart_title_runs_are_one_unit():
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Pt
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    cd = CategoryChartData()
    cd.categories = ["一月", "二月"]
    cd.add_series("销量", (1000, 1200))
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(50), Pt(50), Pt(300), Pt(200), cd).chart
    p = ch.value_axis.axis_title.text_frame.paragraphs[0]
    for t in ("销量（", "万件", "）"):                          # one title split into runs
        p.add_run().text = t
    path = os.path.join(WORK, "chart.pptx")
    prs.save(path)
    units = [u for u in U.make_units(path) if "#c" in u["id"]]
    srcs = {u["src"] for u in units}
    assert "销量（万件）" in srcs and "万件" not in srcs, srcs
    en = {"销量（万件）": "Sales (10k units)", "一月": "January", "二月": "February", "销量": "Sales"}
    out = os.path.join(WORK, "chart_en.pptx")
    U.apply(path, [dict(u, full=en.get(u["src"], u["src"])) for u in units], out)
    ch2 = next(sh for sh in Presentation(out).slides[0].shapes if sh.has_chart).chart
    assert ch2.value_axis.axis_title.text_frame.text == "Sales (10k units)"
    import io, openpyxl                                         # "Edit Data" shows the same words
    wb = openpyxl.load_workbook(io.BytesIO(ch2.part.chart_workbook.xlsx_part.blob))
    cells = {c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str)}
    assert {"January", "February"} <= cells and "一月" not in cells, cells
    # a chart whose data sheet is linked from outside the file (seen in the corpus) must not stop the deck
    prs = Presentation(path)
    part = next(sh for sh in prs.slides[0].shapes if sh.has_chart).chart.part
    ext = part._element.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}externalData")
    rid = part.rels.get_or_add_ext_rel(
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject", "file:///C:/data/book.xlsx")
    ext.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", rid)
    linked = os.path.join(WORK, "chart_linked.pptx")
    prs.save(linked)
    U.apply(linked, [dict(u, full=en.get(u["src"], u["src"])) for u in units], os.path.join(WORK, "chart_linked_en.pptx"))


def _para(p, parts):
    """Fill an <a:p> with runs and fields: parts like ["第", ("fld", "1"), "页"] or [("b", "3倍")]."""
    from lxml import etree
    for x in parts:
        if isinstance(x, str):
            r = etree.SubElement(p, U.qa("r"))
        elif x[0] == "fld":
            r = etree.SubElement(p, U.qa("fld"), id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", type="slidenum")
        else:
            r = etree.SubElement(p, U.qa("r"))
            etree.SubElement(r, U.qa("rPr"), b="1")
        etree.SubElement(r, U.qa("t")).text = x if isinstance(x, str) else x[1]


def test_layout_text_fields_numbers_and_mixed_runs():
    from pptx import Presentation
    from pptx.util import Pt
    prs = Presentation()
    lay = prs.slide_layouts[6]
    s = prs.slides.add_slide(lay)
    tb = s.shapes.add_textbox(Pt(10), Pt(10), Pt(200), Pt(30))
    tb.text_frame.text = "公司口号"
    lay.shapes._spTree.append(copy.deepcopy(tb._element))     # a slogan that lives in the layout
    s.shapes._spTree.remove(tb._element)
    tx = s.shapes.add_textbox(Pt(10), Pt(100), Pt(300), Pt(90)).text_frame._txBody
    for p in tx.findall(U.qa("p")):
        tx.remove(p)
    from lxml import etree
    for parts in (["第", ("fld", "1"), "页"], ["12345"], ["利润增长", ("b", "3倍")]):
        _para(etree.SubElement(tx, U.qa("p")), parts)
    path = os.path.join(WORK, "inherited.pptx")
    prs.save(path)
    units = U.make_units(path)
    by_src = {u["src"]: u for u in units}
    assert by_src["公司口号"]["id"].startswith("L"), units
    assert "第{F0}页" in by_src and "12345" not in by_src and "1" not in by_src, units   # no unit without words
    assert by_src["利润增长3倍"].get("mixed") and not by_src["第{F0}页"].get("mixed")
    assert any(v["code"] == "FIELD" for v in T.check_unit(by_src["第{F0}页"], "Page"))
    en = {"公司口号": "Company motto", "第{F0}页": "Page {F0}", "利润增长3倍": "Profit up 3x"}
    out = os.path.join(WORK, "inherited_en.pptx")
    U.apply(path, [dict(u, full=en[u["src"]]) for u in units], out)
    prs2 = Presentation(out)
    assert "Company motto" in [sh.text_frame.text for sh in U.inherited_index(prs2).values() if sh.has_text_frame]
    p0 = next(sh for sh in prs2.slides[0].shapes if sh.has_text_frame).text_frame._txBody.findall(U.qa("p"))[0]
    assert U._para_text(p0, tokens=True) == "Page {F0}" and p0.find(U.qa("fld")) is not None
    d = Deck(path)                                             # the fitter edits layout text under its per-slide id
    alias = next(k for k in d.idx if k.startswith("s1_L"))
    d.scale_font(alias, 0.5, 18)
    shrunk = os.path.join(WORK, "inherited_shrunk.pptx")
    d.save(shrunk)
    lay = next(sh for sh in U.inherited_index(Presentation(shrunk)).values() if sh.has_text_frame and sh.text_frame.text == "公司口号")
    assert {r.font.size.pt for p in lay.text_frame.paragraphs for r in p.runs} == {9.0}
