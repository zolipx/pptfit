"""Render a pptx with LibreOffice and read back where every word landed.

Rendering is the oracle: every layout decision is judged on the rendered result,
never on estimates. Results are cached by file content, so re-checking an
unchanged file is free.

    pdf = to_pdf("deck.pptx")                 # whole deck
    one = single_slide("deck.pptx", 4, "s4.pptx")
    words = page_words(pdf, 3, slide_w, slide_h)   # 0-based page index
"""
import hashlib, os, shutil, subprocess, tempfile

import pymupdf as fitz
from pptx import Presentation

CACHE = os.environ.get("PPTFIT_CACHE", "/tmp/pptfit_cache")
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def sha(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def to_pdf(pptx_path, timeout=240):
    """Convert with LibreOffice; return the cached pdf path."""
    os.makedirs(CACHE, exist_ok=True)
    out = os.path.join(CACHE, sha(pptx_path)[:24] + ".pdf")
    if os.path.exists(out):
        return out
    work = tempfile.mkdtemp(prefix="lo_")
    src = os.path.join(work, "in.pptx")
    shutil.copy(pptx_path, src)
    profile = "file://" + os.path.join(work, "profile")  # private profile: parallel-safe
    subprocess.run(["soffice", "-env:UserInstallation=" + profile, "--headless",
                    "--convert-to", "pdf", "--outdir", work, src],
                   check=True, timeout=timeout, capture_output=True)
    shutil.move(os.path.join(work, "in.pdf"), out)
    shutil.rmtree(work, ignore_errors=True)
    return out


def single_slide(pptx_path, slide_no, dest):
    """Copy of the deck that keeps only slide `slide_no` (1-based)."""
    prs = Presentation(pptx_path)
    ids = prs.slides._sldIdLst
    for i in sorted((i for i in range(len(ids)) if i != slide_no - 1), reverse=True):
        rid = ids[i].get(R_NS)
        prs.part.drop_rel(rid)
        ids.remove(ids[i])
    prs.save(dest)
    return dest


def render_slide(pptx_path, slide_no):
    """PDF of one slide (page index 0 of the returned pdf)."""
    tmp = os.path.join(tempfile.mkdtemp(prefix="one_"), "s.pptx")
    single_slide(pptx_path, slide_no, tmp)
    pdf = to_pdf(tmp)
    shutil.rmtree(os.path.dirname(tmp), ignore_errors=True)
    return pdf


def page_words(pdf_path, page_index, slide_w, slide_h):
    """Words on a page in slide points: {box, text, size, block}.

    `size` is the rendered font size of the span holding the word, so shrink-on-
    overflow and other render-time effects are visible. `block` groups words of
    one text frame, which keeps overflowing lines attached to their own frame.
    """
    doc = fitz.open(pdf_path)
    page = doc[page_index]
    sx, sy = slide_w / page.rect.width, slide_h / page.rect.height
    spans = []
    for b in page.get_text("dict")["blocks"]:
        for line in b.get("lines", []):
            for s in line["spans"]:
                spans.append((fitz.Rect(s["bbox"]), s["size"]))
    words = []
    for x0, y0, x1, y1, text, block, line, _ in page.get_text("words"):
        r = fitz.Rect(x0, y0, x1, y1)
        cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
        size = next((sz for sr, sz in spans if sr.contains(fitz.Point(cx, cy))), None)
        words.append({"box": [x0 * sx, y0 * sy, (x1 - x0) * sx, (y1 - y0) * sy], "text": text,
                      "size": round(size * sy, 2) if size else None, "block": (block, line)})
    doc.close()
    return words


def page_png(pdf_path, page_index, out_png, dpi=96):
    doc = fitz.open(pdf_path)
    doc[page_index].get_pixmap(dpi=dpi).save(out_png)
    doc.close()
    return out_png


def background_only(pptx_path, dest):
    """Copy with every slide shape removed: what remains is background, master and layout."""
    prs = Presentation(pptx_path)
    for slide in prs.slides:
        tree = slide.shapes._spTree
        for el in list(tree):
            if el.tag.split("}")[-1] in ("sp", "pic", "grpSp", "graphicFrame", "cxnSp", "contentPart"):
                tree.remove(el)
    prs.save(dest)
    return dest


def page_rgb(pdf_path, page_index, dpi):
    import numpy as np
    doc = fitz.open(pdf_path)
    pix = doc[page_index].get_pixmap(dpi=dpi, alpha=False)
    arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)[:, :, :3].copy()
    doc.close()
    return arr
