"""Edits on a pptx, addressed by element id (see extract.py). All lengths in pt.

    d = Deck("in.pptx")
    d.scale_font("s4_12", 0.9, base_size=18)   # relative to the original size
    d.resize("s4_12", dw=0, dh=8, ax="center", ay="center")
    d.move("s4_12", 0, -6); d.trim_line("s4_3", "left", 40); d.set_insets("s4_12", 1.8)
    d.autofit_off(); d.save("out.pptx")

Children of groups are converted to the group's own coordinate space, so callers
always think in slide points. Extension point: add a method here, then offer it
as a candidate move in fit.py.
"""
import re

from lxml import etree
from pptx import Presentation

from .units import shape_index

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
EMU = 12700
qa = lambda t: "{%s}%s" % (A, t)

# Callouts point at something with a tail whose adjust values are fractions of the shape's own size,
# so a plain resize or move drags the tip along. Defaults from the OOXML preset geometry definitions.
WEDGE = ("wedgeRectCallout", "wedgeRoundRectCallout", "wedgeEllipseCallout", "cloudCallout")
WEDGE_DEFAULT = {"adj1": -20833, "adj2": 62500}  # tip offset from the centre, 1/100000 of width, height
LINE_CALLOUT = re.compile(r"(?:borderC|c|accentC|accentBorderC)allout([123])")
LINE_DEFAULT = {1: [18750, -8333, 112500, -38333], 2: [18750, -8333, 18750, -16667, 112500, -46667],
                3: [18750, -8333, 18750, -16667, 100000, -16667, 112963, -8333]}  # (y, x) pairs from the top-left


class Deck:
    def __init__(self, path):
        self.prs = Presentation(path)
        self.idx = shape_index(self.prs)

    def save(self, path):
        self.prs.save(path)

    def _scale(self, eid):
        """slide pt per raw EMU of this shape's own coordinate space (groups shrink/stretch)."""
        sx = sy = 1.0 / EMU
        parts = eid.split("_", 1)[1].split("/")
        prefix = eid.split("_", 1)[0] + "_"
        for i in range(len(parts) - 1):
            g = self.idx[prefix + "/".join(parts[:i + 1])]._element
            x = g.find(qp("grpSpPr")).find(qa("xfrm"))
            ext, chext = x.find(qa("ext")), x.find(qa("chExt"))
            if ext is not None and chext is not None and int(chext.get("cx")) and int(chext.get("cy")):
                sx *= int(ext.get("cx")) / int(chext.get("cx"))
                sy *= int(ext.get("cy")) / int(chext.get("cy"))
        return sx, sy

    def _xfrm(self, eid):
        el = self.idx[eid]._element
        sp = el.find(qp("spPr")) if el.find(qp("spPr")) is not None else el.find(qp("grpSpPr"))
        if sp is None:  # graphicFrame
            return el.find(qp("xfrm")), el
        return sp.find(qa("xfrm")), el

    def geometry(self, eid):
        x, _ = self._xfrm(eid)
        if x is None or x.find(qa("off")) is None:  # placeholder inheriting its geometry from the layout
            sh = self.idx[eid]
            left, top, width, height = sh.left, sh.top, sh.width, sh.height
            sh.left, sh.top, sh.width, sh.height = left, top, width, height  # writes an explicit xfrm
            x, _ = self._xfrm(eid)
        off, ext = x.find(qa("off")), x.find(qa("ext"))
        return [int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))]

    def _set(self, eid, g):
        x, _ = self._xfrm(eid)
        x.find(qa("off")).set("x", str(int(round(g[0])))); x.find(qa("off")).set("y", str(int(round(g[1]))))
        x.find(qa("ext")).set("cx", str(max(int(round(g[2])), 1))); x.find(qa("ext")).set("cy", str(max(int(round(g[3])), 1)))

    def _callout(self, eid):
        """(prst, avLst, adjust names, values) of a callout shape, else None."""
        g = self.idx[eid]._element.find(".//" + qa("prstGeom"))
        prst = g.get("prst", "") if g is not None else ""
        m = LINE_CALLOUT.fullmatch(prst)
        if prst in WEDGE:
            names, vals = ["adj1", "adj2"], dict(WEDGE_DEFAULT)
        elif m:
            n = int(m.group(1))
            names = ["adj%d" % (i + 1) for i in range(2 * n + 2)]
            vals = dict(zip(names, LINE_DEFAULT[n]))
        else:
            return None
        av = g.find(qa("avLst"))
        if av is None:
            av = etree.SubElement(g, qa("avLst"))
        for gd in av.findall(qa("gd")):
            f = (gd.get("fmla") or "").split()
            if len(f) == 2 and f[0] == "val":
                vals[gd.get("name")] = int(f[1])
        return prst, av, names, vals

    def _place(self, eid, old, new):
        """Set new geometry; a callout's tip keeps its absolute place (the line's start stays on the box)."""
        c = self._callout(eid)
        self._set(eid, new)
        if not c:
            return
        prst, av, names, vals = c
        ox, oy, ow, oh = old
        nx, ny, nw, nh = new[0], new[1], max(new[2], 1), max(new[3], 1)
        if prst in WEDGE:
            px, py = ox + ow / 2 + vals["adj1"] * ow / 1e5, oy + oh / 2 + vals["adj2"] * oh / 1e5
            vals["adj1"], vals["adj2"] = (px - nx - nw / 2) / nw * 1e5, (py - ny - nh / 2) / nh * 1e5
        else:
            for k in range(1, len(names) // 2):
                ya, xa = names[2 * k], names[2 * k + 1]
                py, px = oy + vals[ya] * oh / 1e5, ox + vals[xa] * ow / 1e5
                vals[ya], vals[xa] = (py - ny) / nh * 1e5, (px - nx) / nw * 1e5
        have = {gd.get("name"): gd for gd in av.findall(qa("gd"))}
        for n in names:
            gd = have.get(n)
            if gd is None:
                gd = etree.SubElement(av, qa("gd"), name=n)
            gd.set("fmla", "val %d" % round(vals[n]))

    def resize(self, eid, dw=0.0, dh=0.0, ax="left", ay="top"):
        """Grow (or shrink) by dw/dh pt, keeping the anchor side or the centre fixed."""
        sx, sy = self._scale(eid)
        g = self.geometry(eid)
        old = list(g)
        rw, rh = dw / sx, dh / sy
        g[0] -= {"left": 0, "center": rw / 2, "right": rw}[ax]
        g[1] -= {"top": 0, "center": rh / 2, "bottom": rh}[ay]
        g[2] += rw; g[3] += rh
        self._place(eid, old, g)

    def move(self, eid, dx=0.0, dy=0.0):
        sx, sy = self._scale(eid)
        g = self.geometry(eid)
        old = list(g)
        g[0] += dx / sx; g[1] += dy / sy
        self._place(eid, old, g)

    def trim_line(self, eid, end, d):
        """Shorten a horizontal/vertical line from one end by d pt."""
        sx, sy = self._scale(eid)
        g = self.geometry(eid)
        if g[2] >= g[3]:
            g[2] -= d / sx
            if end == "left":
                g[0] += d / sx
        else:
            g[3] -= d / sy
            if end == "top":
                g[1] += d / sy
        self._set(eid, g)

    def _bodies(self, eid):
        el = self.idx[eid]._element
        return el.iter(qp("txBody")) if el.find(".//" + qp("txBody")) is not None else el.iter(qa("txBody"))

    def scale_font(self, eid, factor, base_size):
        """Set every run to `factor` x its current size (call on a fresh base deck,
        so "current" is the original). Runs without a size get base_size x factor."""
        for tx in self._bodies(eid):
            for node in tx.iter(qa("rPr"), qa("endParaRPr"), qa("defRPr")):
                sz = int(node.get("sz") or base_size * 100)
                node.set("sz", str(max(int(round(sz * factor / 50.0)) * 50, 100)))  # half-point steps
            for r in tx.iter(qa("r")):
                if r.find(qa("rPr")) is None:
                    r.insert(0, etree.Element(qa("rPr"), sz=str(max(int(round(base_size * factor * 2)) * 50, 100))))

    def line_spacing(self, eid, factor, base_size=18.0):
        """Scale line spacing and space before/after (on a fresh base deck). Many CJK decks use
        1.5x line spacing, which is the first thing to give; it never goes below 1.1x the font
        size, the usual floor for readable body text."""
        for tx in self._bodies(eid):
            for p in tx.iter(qa("p")):
                sizes = [int(r.get("sz")) for r in p.iter(qa("rPr")) if r.get("sz")] or [int(base_size * 100)]
                for tag in ("lnSpc", "spcBef", "spcAft"):
                    node = p.find(qa("pPr") + "/" + qa(tag)) if p.find(qa("pPr")) is not None else None
                    if node is None:
                        continue
                    pct, pts = node.find(qa("spcPct")), node.find(qa("spcPts"))
                    if pct is not None:
                        old = int(pct.get("val"))
                        floor = min(110000, old) if tag == "lnSpc" else 0
                        pct.set("val", str(max(int(old * factor), floor)))
                    if pts is not None:
                        old = int(pts.get("val"))
                        floor = min(int(1.1 * max(sizes)), old) if tag == "lnSpc" else 0
                        pts.set("val", str(max(int(old * factor), floor)))

    def has_spacing(self, eid):
        return any(True for tx in self._bodies(eid) for _ in tx.iter(qa("lnSpc"), qa("spcBef"), qa("spcAft")))

    def set_insets(self, eid, pt=1.8):
        for tx in self._bodies(eid):
            bp = tx.find(qa("bodyPr"))
            if bp is not None:
                for k in ("lIns", "rIns", "tIns", "bIns"):
                    bp.set(k, str(int(pt * EMU)))

    def latin_font(self, name="Arial"):
        """Give translated text a real Latin font where the deck uses a CJK font for Latin
        (SimSun, Microsoft YaHei, ...): those render English badly. Latin fonts stay as they are."""
        n = 0
        theme_cjk = self._theme_latin_is_cjk()
        for eid in self.idx:
            for tx in self._bodies(eid):
                for node in tx.iter(qa("rPr"), qa("endParaRPr"), qa("defRPr")):
                    lat = node.find(qa("latin"))
                    if lat is not None and is_cjk_font(lat.get("typeface") or ""):
                        lat.set("typeface", name); n += 1
                    elif lat is None and theme_cjk:
                        node.insert(0, etree.Element(qa("latin"), typeface=name)); n += 1
        return n

    def _theme_latin_is_cjk(self):
        for part in self.prs.part.package.iter_parts():
            if str(part.partname).startswith("/ppt/theme/"):
                root = etree.fromstring(part.blob)
                for node in root.iter(qa("minorFont"), qa("majorFont")):
                    lat = node.find(qa("latin"))
                    if lat is not None and is_cjk_font(lat.get("typeface") or ""):
                        return True
                return False
        return False

    def add_note(self, slide_no, lines):
        """Append lines to the speaker notes of a slide (1-based)."""
        tf = self.prs.slides[slide_no - 1].notes_slide.notes_text_frame
        if tf is None:
            raise ValueError("slide %d: its notes page has no text body, so notes cannot be added" % slide_no)
        for line in lines:
            tf.add_paragraph().text = line

    def autofit_off(self):
        """Drop shrink-on-overflow everywhere: sizes must be what we set, not what a renderer decides."""
        n = 0
        for eid in self.idx:
            for tx in self._bodies(eid):
                bp = tx.find(qa("bodyPr"))
                na = bp.find(qa("normAutofit")) if bp is not None else None
                if na is not None:
                    bp.remove(na)
                    etree.SubElement(bp, qa("noAutofit"))
                    n += 1
        return n


CJK_FONT_HINTS = ("yahei", "simsun", "simhei", "nsimsun", "dengxian", "kaiti", "fangsong", "heiti", "songti", "stxihei",
                  "pingfang", "hiragino", "mincho", "gothic", "meiryo", "malgun", "batang", "gulim", "cjk", "source han", "msyh")


def is_cjk_font(name):
    return any(ord(c) > 127 for c in name) or any(h in name.lower() for h in CJK_FONT_HINTS)


def qp(t):
    return "{%s}%s" % (P, t)
