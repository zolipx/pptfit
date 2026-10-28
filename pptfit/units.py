"""Translation units: split any deck into pieces a translator fills, and write them back.

A unit is one paragraph of a shape or table cell, one distinct string of a chart
or SmartArt, or one paragraph of the speaker notes. Ids are derived from the
file, so the same code works on any deck:

    s4_12#p0          paragraph 0 of shape 12 on slide 4 (groups: s4_7/12#p0)
    s5_3/r1c2#p0      table cell row 1, col 2
    s9_8#c3           chart string 3: a title/label paragraph or a category/series name (matched by its source text)
    s2_5#d0           SmartArt paragraph 0 (matched by its source text)
    s4#notes#p1       notes paragraph 1
    L3_7#p0           paragraph 0 of shape 7 in slide layout 3 (M1_4#p0: slide master 1);
                      only non-placeholder shapes of layouts that slides use: logos, slogans, footers

A field (slide number, date) inside a paragraph shows up as {F0}, {F1}, ... in `src`; the
translation keeps the token and the live field is put back there. Paragraphs with no words
(numbers, fields only) are not units. `mixed: true` marks a paragraph whose runs differ in
formatting (one bold or coloured word); it is written back in the first run's format and
the fitter notes it for the reviewer.

Translations file: [{id, src, full, short, concepts:[{zh,en}], back}], written by
the model. `apply()` writes the chosen variant ("full" or "short") per unit.
"""
import copy, re

from lxml import etree
from pptx import Presentation

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
DGM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
R_DM = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}dm"
qa = lambda t: "{%s}%s" % (A, t)


def shape_index(prs):
    """{element id: shape} for every shape, groups expanded, ids as in extract.py."""
    out = {}

    def walk(shapes, prefix):
        for sh in shapes:
            ident = prefix + str(sh.shape_id)
            out[ident] = sh
            try:
                is_group = sh.shape_type == 6
            except NotImplementedError:
                is_group = False
            if is_group:
                walk(sh.shapes, ident + "/")

    for i, slide in enumerate(prs.slides, 1):
        walk(slide.shapes, "s%d_" % i)
    return out


FIELD = re.compile(r"\{F(\d+)\}")


def _para_text(p, tokens=False):
    parts, k = [], 0
    for node in p:
        tag = node.tag.split("}")[-1]
        if tag == "fld" and tokens:
            parts.append("{F%d}" % k)
            k += 1
        elif tag in ("r", "fld"):
            t = node.find(qa("t"))
            parts.append(t.text or "" if t is not None else "")
        elif tag == "br":
            parts.append("\n")
    return "".join(parts)


def has_words(text):
    return re.search(r"[^\W\d_]", FIELD.sub("", text)) is not None


def _mixed(p):
    """True if the runs with text differ in visible formatting (bold, italic, size, colour, ...)."""
    keys = set()
    for r in p.findall(qa("r")):
        t = r.find(qa("t"))
        if t is None or not (t.text or "").strip():
            continue
        rpr = r.find(qa("rPr"))
        if rpr is None:
            keys.add(None)
            continue
        attrs = tuple(rpr.get(a) for a in ("b", "i", "u", "strike", "sz", "baseline"))
        fills = tuple(etree.tostring(c) for c in rpr if c.tag.split("}")[-1] in ("solidFill", "highlight", "gradFill"))
        keys.add((attrs, fills))
    return len(keys) > 1


def inherited_index(prs):
    """{id: shape} for the non-placeholder shapes of the layouts slides use and of their masters."""
    out = {}

    def walk(shapes, prefix):
        for sh in shapes:
            if sh.is_placeholder:
                continue
            ident = prefix + str(sh.shape_id)
            out[ident] = sh
            try:
                is_group = sh.shape_type == 6
            except NotImplementedError:
                is_group = False
            if is_group:
                walk(sh.shapes, ident + "/")

    num = lambda part: re.search(r"(\d+)\.xml$", str(part.partname)).group(1)
    for slide in prs.slides:
        lay = slide.slide_layout
        walk(lay.shapes, "L%s_" % num(lay.part))
        walk(lay.slide_master.shapes, "M%s_" % num(lay.slide_master.part))
    return out


def _diagram_parts(prs):
    return [p for p in prs.part.package.iter_parts() if str(p.partname).startswith("/ppt/diagrams/")]


def make_units(pptx_path):
    prs = Presentation(pptx_path)
    units = []

    def add_txbody(eid, tx):
        if tx is None:
            return
        for k, p in enumerate(tx.findall(qa("p"))):
            text = _para_text(p, tokens=True)
            if has_words(text):
                units.append({"id": "%s#p%d" % (eid, k), "src": text, **({"mixed": True} if _mixed(p) else {})})

    for eid, sh in shape_index(prs).items():
        el = sh._element
        tx = el.find("{http://schemas.openxmlformats.org/presentationml/2006/main}txBody")
        add_txbody(eid, tx)
        if getattr(sh, "has_table", False) and sh.has_table:
            for r, row in enumerate(sh.table.rows):
                for c, cell in enumerate(row.cells):
                    add_txbody("%s/r%dc%d" % (eid, r, c), cell._tc.find(qa("txBody")))
        if getattr(sh, "has_chart", False) and sh.has_chart:
            seen = []
            for s, _ in _texts(sh.chart._chartSpace, chart=True):
                if s not in seen and not re.fullmatch(r"[-\d.,%\s]+", s):
                    seen.append(s)
            units += [{"id": "%s#c%d" % (eid, k), "src": t} for k, t in enumerate(seen)]
        rel = el.find(".//{%s}relIds" % DGM)
        if rel is not None:  # SmartArt: texts live in the diagram data part
            data = sh.part.related_part(rel.get(R_DM))
            seen = []
            for s, _ in _texts(etree.fromstring(data.blob)):
                if s not in seen:
                    seen.append(s)
            units += [{"id": "%s#d%d" % (eid, k), "src": t} for k, t in enumerate(seen)]
    for eid, sh in inherited_index(prs).items():
        add_txbody(eid, sh._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}txBody"))
    for i, slide in enumerate(prs.slides, 1):
        if slide.has_notes_slide:
            add_txbody("s%d#notes" % i, slide.notes_slide.notes_text_frame._txBody)
    return units


def _texts(root, chart=False):
    """[(source text, writer)] for the strings under root: each rich-text paragraph as a whole (one title
    is often split into runs, e.g. "销量（" "万件" "）"; translating runs apart garbles it), and for
    charts the cached category and series names."""
    out = []
    for p in root.iter(qa("p")):
        s = _para_text(p)
        if s.strip():
            out.append((s, lambda text, p=p: set_paragraph_text(p, text)))
    if chart:
        for v in root.iter("{%s}v" % C):
            anc = [a.tag.split("}")[-1] for a in v.iterancestors()]
            if any(a.startswith("num") for a in anc[:3]) or not (v.text and v.text.strip()):
                continue
            out.append((v.text, lambda text, v=v: setattr(v, "text", text)))
    return out


def set_paragraph_text(p, text):
    """Replace a paragraph's text, keeping the first run's formatting; '\n' becomes a line break,
    {F0}, {F1}, ... put back the paragraph's own fields (slide number, date)."""
    runs = [n for n in p if n.tag.split("}")[-1] in ("r", "fld")]
    flds = [copy.deepcopy(n) for n in p if n.tag.split("}")[-1] == "fld"]
    rpr = None
    if runs and runs[0].find(qa("rPr")) is not None:
        rpr = copy.deepcopy(runs[0].find(qa("rPr")))
    elif p.find(qa("endParaRPr")) is not None:
        rpr = copy.deepcopy(p.find(qa("endParaRPr")))
        rpr.tag = qa("rPr")
    for n in list(p):
        if n.tag.split("}")[-1] in ("r", "fld", "br"):
            p.remove(n)
    end = p.find(qa("endParaRPr"))
    pieces = text.split("\n")
    nodes = []
    for i, piece in enumerate(pieces):
        if i:
            br = etree.Element(qa("br"))
            if rpr is not None:
                br.append(copy.deepcopy(rpr))
            nodes.append(br)
        segs = [s for s in FIELD.split(piece)]          # text, field number, text, ...
        for j, seg in enumerate(segs):
            if j % 2 and int(seg) < len(flds):
                nodes.append(copy.deepcopy(flds[int(seg)]))
                continue
            if j % 2 or (not seg and len(segs) > 1):
                continue
            r = etree.Element(qa("r"))
            if rpr is not None:
                r.append(copy.deepcopy(rpr))
            t = etree.SubElement(r, qa("t"))
            t.text = seg
            nodes.append(r)
    for n in nodes:
        if end is not None:
            end.addprevious(n)
        else:
            p.append(n)


def _update_workbook(chart, by_src):
    """PowerPoint's "Edit Data" opens the chart's embedded workbook: give it the same strings,
    or the labels fall back to the source language as soon as someone edits the data."""
    import io, sys, zipfile
    import openpyxl
    from openpyxl.utils.exceptions import InvalidFileException
    ext = chart._chartSpace.find("{%s}externalData" % C)
    if ext is None:
        return
    rel = chart.part.rels[ext.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")]
    if rel.is_external:  # a workbook linked from outside the file: nothing embedded to edit
        return
    xlsx = rel.target_part
    try:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx.blob))
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError) as e:  # an odd workbook must not stop the deck
        print("chart data sheet not translated (%s): %s" % (xlsx.partname, e), file=sys.stderr)
        return
    hit = False
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value in by_src:
                    c.value, hit = by_src[c.value], True
    if hit:
        buf = io.BytesIO()
        wb.save(buf)
        chart.part.chart_workbook.update_from_xlsx_blob(buf.getvalue())


def apply(pptx_in, translations, pptx_out, choice=None, default="full"):
    """Write translations into a copy. choice: {unit id: "full"|"short"|...}."""
    choice = choice or {}
    prs = Presentation(pptx_in)
    idx = shape_index(prs)
    idx.update(inherited_index(prs))
    chosen = {}
    for u in translations:
        v = choice.get(u["id"], default)
        chosen[u["id"]] = u.get(v) or u.get("full") or u["src"]
    by_src = {}  # chart / SmartArt strings are matched by their source text
    for u in translations:
        uid = u["id"]
        if "#c" in uid or "#d" in uid:
            by_src[u["src"]] = chosen[uid]
            continue
        eid, para = uid.rsplit("#p", 1)
        if "#notes" in eid:
            slide = prs.slides[int(eid.split("#")[0][1:]) - 1]
            tx = slide.notes_slide.notes_text_frame._txBody
        else:
            m = re.match(r"(.*)/r(\d+)c(\d+)$", eid)
            if m:
                tx = idx[m.group(1)].table.cell(int(m.group(2)), int(m.group(3)))._tc.find(qa("txBody"))
            else:
                tx = idx[eid]._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}txBody")
        set_paragraph_text(tx.findall(qa("p"))[int(para)], chosen[uid])
    if by_src:
        for sh in idx.values():
            if getattr(sh, "has_chart", False) and sh.has_chart:
                for s, write in _texts(sh.chart._chartSpace, chart=True):
                    if s in by_src:
                        write(by_src[s])
                _update_workbook(sh.chart, by_src)
        for part in _diagram_parts(prs):
            root = etree.fromstring(part.blob)
            hit = False
            for s, write in _texts(root):
                if s in by_src:
                    write(by_src[s])
                    hit = True
            if hit:
                part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    prs.save(pptx_out)
    return chosen
