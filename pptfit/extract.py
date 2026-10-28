#!/usr/bin/env python3
"""Extract pptx structure -> source.json."""
import json, sys, re
from pptx import Presentation
from lxml import etree

EMU = 12700.0
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
C = "http://schemas.openxmlformats.org/drawingml/2006/chart"

def qa(tag):
    return "{%s}%s" % (A, tag)

def qp(tag):
    return "{%s}%s" % (P, tag)

def emu(v):
    return float(v) / EMU

def color_of(node):
    if node is None:
        return None
    s = node.find(".//" + qa("srgbClr"))
    if s is not None and s.get("val"):
        return s.get("val").upper()
    sch = node.find(".//" + qa("schemeClr"))
    if sch is not None and sch.get("val"):
        return "scheme:" + sch.get("val")
    return "other"

def sppr(el):
    for ch in el:
        if ch.tag in (qp("spPr"), qp("grpSpPr")):
            return ch
    return None

def xfrm(el):
    sp = sppr(el)
    node = sp.find(qa("xfrm")) if sp is not None else None
    if node is None:
        return None
    def pair(n, a, b):
        return None if n is None else (int(n.get(a) or 0), int(n.get(b) or 0))
    return pair(node.find(qa("off")), "x", "y"), pair(node.find(qa("ext")), "cx", "cy"), pair(node.find(qa("chOff")), "x", "y"), pair(node.find(qa("chExt")), "cx", "cy")

def prst_of(el):
    sp = sppr(el)
    if sp is None:
        return None
    pg = sp.find(qa("prstGeom"))
    if pg is not None:
        return pg.get("prst")
    if sp.find(qa("custGeom")) is not None:
        return "custGeom"
    return None

FAINT = 20000  # alpha (1/1000 %) at or below which a fill or line is a tint, not an obstacle: decorative bands behind text


def faint(node):
    """True if every colour in node is drawn at <= 20% opacity."""
    alphas = [int(a.get("val", "100000")) for a in node.iter(qa("alpha"))]
    return bool(alphas) and max(alphas) <= FAINT


def fill_line(el):
    sp = sppr(el)
    fill = line = None
    vis_f = vis_l = False
    if sp is not None:
        if sp.find(qa("noFill")) is None:
            for tag, key in (("solidFill", None), ("gradFill", "grad"), ("blipFill", "blip"), ("pattFill", "patt")):
                n = sp.find(qa(tag))
                if n is not None:
                    vis_f = not faint(n)
                    fill = key or color_of(n)
                    break
        ln = sp.find(qa("ln"))
        if ln is not None and ln.find(qa("noFill")) is None:
            sf = ln.find(qa("solidFill"))
            if (sf is not None and not faint(sf)) or (sf is None and ln.get("w") and int(ln.get("w") or 0) > 0):
                vis_l = True
                line = color_of(sf) if sf is not None else "default"
    style = el.find(qp("style"))
    if style is not None:
        if not vis_f and (sp is None or sp.find(qa("noFill")) is None):
            fr = style.find(qa("fillRef"))
            if fr is not None and fr.get("idx") not in (None, "0"):
                vis_f, fill = True, "style:" + fr.get("idx")
        explicit_no = sp is not None and sp.find(qa("ln")) is not None and sp.find(qa("ln")).find(qa("noFill")) is not None
        if not vis_l and not explicit_no:
            lr = style.find(qa("lnRef"))
            if lr is not None and lr.get("idx") not in (None, "0"):
                vis_l, line = True, line or ("style:" + lr.get("idx"))
    return fill, vis_f, line, vis_l

def theme_fonts(prs):
    out = {"+mj-lt": "Calibri", "+mn-lt": "Calibri", "+mj-ea": "", "+mn-ea": ""}
    blob = None
    for part in prs.part.package.iter_parts():
        if str(part.partname).startswith("/ppt/theme/"):
            blob = part.blob
            break
    if not blob:
        return out
    root = etree.fromstring(blob)
    for scheme, key in (("majorFont", "+mj"), ("minorFont", "+mn")):
        node = root.find(".//" + qa(scheme))
        if node is None:
            continue
        lat, ea = node.find(qa("latin")), node.find(qa("ea"))
        if lat is not None and lat.get("typeface"):
            out[key + "-lt"] = lat.get("typeface")
        if ea is not None and ea.get("typeface"):
            out[key + "-ea"] = ea.get("typeface")
    return out

def face(val, theme, which):
    if not val:
        return None
    if val.startswith("+"):
        return theme.get(val) or theme.get("+mn-" + which) or val
    return val

def rpr_sz(rPr):
    if rPr is not None and rPr.get("sz"):
        return float(rPr.get("sz")) / 100.0
    return None

def rpr_fonts(rPr, theme):
    if rPr is None:
        return None, None
    lat, ea = rPr.find(qa("latin")), rPr.find(qa("ea"))
    return face(lat.get("typeface") if lat is not None else None, theme, "lt"), face(ea.get("typeface") if ea is not None else None, theme, "ea")

def body_meta(tx):
    bp = tx.find(qa("bodyPr")) if tx is not None else None
    insets = {"l": 7.2, "r": 7.2, "t": 3.6, "b": 3.6}
    wrap, anchor, scale, autofit = True, "top", 1.0, None
    if bp is None:
        return insets, wrap, anchor, scale, autofit
    if bp.get("wrap") == "none":
        wrap = False
    anchor = {"ctr": "center", "b": "bottom", "t": "top"}.get(bp.get("anchor") or "t", "top")
    for k, attr in (("l", "lIns"), ("r", "rIns"), ("t", "tIns"), ("b", "bIns")):
        if bp.get(attr):
            insets[k] = int(bp.get(attr)) / EMU
    na = bp.find(qa("normAutofit"))
    if na is not None:
        autofit = "norm"
        if na.get("fontScale"):
            scale = int(na.get("fontScale")) / 100000.0
    elif bp.find(qa("spAutoFit")) is not None:
        autofit = "sp"
    elif bp.find(qa("noAutofit")) is not None:
        autofit = "none"
    return insets, wrap, anchor, scale, autofit

def inherit_sz(rPr, p, tx):
    sz = rpr_sz(rPr)
    if sz:
        return sz, False
    pPr = p.find(qa("pPr")) if p is not None else None
    if pPr is not None:
        sz = rpr_sz(pPr.find(qa("defRPr")))
        if sz:
            return sz, False
    lvl = int(pPr.get("lvl")) if pPr is not None and pPr.get("lvl") else 0
    lst = tx.find(qa("lstStyle")) if tx is not None else None
    if lst is not None:
        node = lst.find(qa("lvl%dpPr" % (lvl + 1)))
        if node is not None:
            sz = rpr_sz(node.find(qa("defRPr")))
            if sz:
                return sz, False
    return 18.0, True

def paragraphs_from(tx, theme, scale):
    paras = []
    if tx is None:
        return paras
    for p in tx.findall(qa("p")):
        pPr = p.find(qa("pPr"))
        algn = {"ctr": "center", "r": "right", "just": "justify", "dist": "distribute"}.get((pPr.get("algn") if pPr is not None else None) or "l", "left")
        lvl = int(pPr.get("lvl")) if pPr is not None and pPr.get("lvl") else 0
        runs = []
        for r in p.findall(qa("r")):
            t = r.find(qa("t"))
            text = (t.text or "") if t is not None else ""
            rPr = r.find(qa("rPr"))
            sz, guess = inherit_sz(rPr, p, tx)
            latin, ea = rpr_fonts(rPr, theme)
            if pPr is not None and (not latin or not ea):
                l2, e2 = rpr_fonts(pPr.find(qa("defRPr")), theme)
                latin = latin or l2
                ea = ea or e2
            runs.append({"text": text, "size": round((sz or 18) * scale, 3), "size_guess": guess, "latin": latin or "Calibri", "ea": ea or ""})
        if not runs:
            end = p.find(qa("endParaRPr"))
            sz, guess = inherit_sz(end, p, tx)
            runs.append({"text": "", "size": round((sz or 18) * scale, 3), "size_guess": guess, "latin": "Calibri", "ea": ""})
        paras.append({"align": algn, "level": lvl, "runs": runs})
    return paras

def plain(paras):
    return "\n".join("".join(r["text"] for r in p["runs"]) for p in paras)

def nums(text):
    return re.findall(r"\d+(?:\.\d+)?%?", text or "")

def txbody(el):
    for ch in el:
        if ch.tag == qp("txBody"):
            return ch
    return None

def kind_of(shape, el, text):
    tag = el.tag.split("}")[-1]
    if tag == "pic":
        return "picture"
    if tag == "cxnSp":
        return "line"
    if tag == "grpSp":
        return "group"
    if tag == "graphicFrame":
        if getattr(shape, "has_table", False):
            return "table"
        if getattr(shape, "has_chart", False):
            return "chart"
        xml = etree.tostring(el, encoding=str)
        return "smartart" if ("diagram" in xml or "dgm:" in xml) else "other"
    prst = prst_of(el)
    try:
        st = int(shape.shape_type) if shape.shape_type is not None else None
    except NotImplementedError:  # python-pptx has no type for some shapes
        st = None
    if prst == "line" or st == 9:
        return "line"
    if text and text.strip():
        return "text"
    return "other"

def chart_strings(shape):
    out = []
    try:
        el = shape.chart._element
    except ValueError:  # a graphic frame without a chart
        return out
    for p in el.iter(qa("p")):  # a paragraph as a whole: titles are often split into runs
        text = "".join(t.text or "" for t in p.iter(qa("t")))
        if text.strip() and text not in out:
            out.append(text)
    for v in el.findall(".//" + "{%s}v" % C):
        parent = v.getparent()
        gp = parent.getparent() if parent is not None else None
        gtag = gp.tag.split("}")[-1] if gp is not None else ""
        if gtag.startswith("num"):
            continue
        if gtag in ("strCache", "strLit", "tx") or (parent is not None and parent.tag.split("}")[-1] in ("tx", "pt")):
            if v.text and v.text.strip() and not re.fullmatch(r"-?\d+(?:\.\d+)?", v.text.strip()):
                if v.text not in out:
                    out.append(v.text)
    return out

def walk(container, prefix, xf, theme, slide_area, out, path_ids):
    for shape in container:
        el = shape._element
        tag = el.tag.split("}")[-1]
        geo = xfrm(el)
        if geo is None or geo[0] is None:
            try:
                off, ext = (int(shape.left), int(shape.top)), (int(shape.width), int(shape.height))
                choff = chext = None
            except TypeError:  # a placeholder with no position of its own or inherited: nothing to place
                continue
        else:
            off, ext, choff, chext = geo
            ext = ext or (0, 0)
        ax, ay, aw, ah = xf(off[0], off[1], ext[0], ext[1])
        sid = int(shape.shape_id)
        ident = prefix + str(sid)
        fill, vis_f, line, vis_l = fill_line(el)
        tx = txbody(el)
        insets, wrap, anchor, scale, autofit = body_meta(tx)
        paras = paragraphs_from(tx, theme, scale) if tx is not None else []
        text = plain(paras)
        kind = kind_of(shape, el, text)
        bbox = [round(emu(ax), 3), round(emu(ay), 3), round(emu(aw), 3), round(emu(ah), 3)]
        rec = {
            "id": ident, "shape_id": sid, "path": path_ids + [sid], "kind": kind, "name": shape.name,
            "bbox": bbox, "background": bbox[2] * bbox[3] >= 0.9 * slide_area, "prst": prst_of(el),
            "fill": fill, "line": line, "visible_fill": bool(vis_f) or kind == "picture", "visible_line": bool(vis_l),
            "placeholder": (el.find(".//" + qp("ph")).get("type") or "body") if el.find(".//" + qp("ph")) is not None else None,
            "is_title": False, "paragraphs": paras, "text": text,
            "insets": {k: round(v, 3) for k, v in insets.items()}, "wrap": wrap, "anchor": anchor,
            "font_scale": scale, "autofit": autofit, "numbers": nums(text),
        }
        if kind == "text" and (rec["placeholder"] in ("title", "ctrTitle") or (
                bbox[1] < 0.12 * (slide_area ** 0.5) and len(text.strip()) <= 60 and "\n" not in text.strip()
                and not rec["visible_fill"] and not rec["visible_line"])):
            rec["is_title"] = True
        if kind == "chart":
            rec["chart_texts"] = chart_strings(shape)
            rec["numbers"] += [n for n in nums(" ".join(rec["chart_texts"])) if n not in rec["numbers"]]
        if kind == "table":
            cells, y = [], ay
            tbl = shape.table
            for ri, row in enumerate(tbl.rows):
                x, rh = ax, int(row.height)
                for ci, cell in enumerate(row.cells):
                    cw = int(tbl.columns[ci].width)
                    c_tx = cell._tc.find(qa("txBody"))
                    cin, _, _, cfs, _ = body_meta(c_tx)
                    cparas = paragraphs_from(c_tx, theme, cfs)
                    ct = plain(cparas)
                    cells.append({"id": "%s/r%dc%d" % (ident, ri, ci), "kind": "text", "bbox": [round(emu(x), 3), round(emu(y), 3), round(emu(cw), 3), round(emu(rh), 3)], "paragraphs": cparas, "text": ct, "insets": cin, "numbers": nums(ct), "parent": ident})
                    x += cw
                y += rh
            rec["cells"] = cells
        out.append(rec)
        if tag == "grpSp":
            if not choff:
                choff = (0, 0)
            if not chext or chext[0] == 0 or chext[1] == 0:
                chext = (ext[0] or 1, ext[1] or 1)
            def child_xf(x, y, w, h, ax=ax, ay=ay, aw=aw, ah=ah, choff=choff, chext=chext):
                sx = aw / chext[0] if chext[0] else 1
                sy = ah / chext[1] if chext[1] else 1
                return (ax + (x - choff[0]) * sx, ay + (y - choff[1]) * sy, w * sx, h * sy)
            walk(shape.shapes, ident + "/", child_xf, theme, slide_area, out, path_ids + [sid])

def extract(path):
    prs = Presentation(path)
    sw, sh = emu(int(prs.slide_width)), emu(int(prs.slide_height))
    theme = theme_fonts(prs)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        shapes = []
        ident = lambda x, y, w, h: (x, y, w, h)
        walk(slide.shapes, "s%d_" % i, ident, theme, sw * sh, shapes, [])
        # Footer bars, logos and rules usually live in the layout/master: read-only obstacles.
        inherited = []
        if slide._element.get("showMasterSp") != "0":
            lay = slide.slide_layout
            num = lambda part: re.search(r"(\d+)\.xml$", str(part.partname)).group(1)
            homes = [("s%d_L" % i, "L%s_" % num(lay.part), lay.shapes)]
            if lay._element.get("showMasterSp") != "0":
                homes.append(("s%d_M" % i, "M%s_" % num(lay.slide_master.part), lay.slide_master.shapes))
            for prefix, home, shapes_ in homes:  # home: the id units.py gives the same shape (L3_19), shared by every slide
                k = len(inherited)
                walk([x for x in shapes_ if not x.is_placeholder], prefix, ident, theme, sw * sh, inherited, [])
                for rec in inherited[k:]:
                    rec["home"] = home + rec["id"][len(prefix):]
        for rec in inherited:
            rec["inherited"] = True
            rec["is_title"] = False
        shapes += inherited
        notes = ""
        if slide.has_notes_slide:
            tf = slide.notes_slide.notes_text_frame  # None when the notes page has no body placeholder
            notes = (tf.text or "") if tf is not None else ""
        slides.append({"index": i, "notes": notes, "shapes": shapes})
    return {"slide_w": round(sw, 3), "slide_h": round(sh, 3), "theme": theme, "slides": slides}

def main():  # python -m pptfit.extract in.pptx out.json
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else "source.json"
    data = extract(src)
    json.dump(data, open(dst, "w"), ensure_ascii=False, indent=1)
    print("slides", len(data["slides"]), "shapes", sum(len(s["shapes"]) for s in data["slides"]), "->", dst)

if __name__ == "__main__":
    main()
